import json
from datetime import datetime
import pytest
from openpyxl import Workbook
from tools.prepare_site import create
from analysis.contracts import Profile


def test_preflight_form_only_embeds_allowlisted_inventory(tmp_path):
    wb=Workbook();wb.remove(wb.active)
    tables={
        '00_Metadata':[['contract_version'],['2.0']],
        '01_Capabilities':[['contract_version','source_name','column_name','available','is_required'],['2.0','Source','Time',0,0]],
        '02_Activities':[['activity_code','activity_name','activity_category'],['CONS','Consultation','Clinical']],
        '03_AppointmentInventory':[['source_state','activity_code','appointment_status','machine','distinct_appointments'],['AVAILABLE','CONS','Complete','M1',25]],
        '04_MachineInventory':[['source_state','machine'],['AVAILABLE','M1']],
        'Private':[['patient_key'],['DO_NOT_EMBED']],
    }
    for title,rows in tables.items():
        sheet=wb.create_sheet(title)
        for row in rows:sheet.append(row)
    path=tmp_path/'preflight.xlsx';wb.save(path)
    output=create(path,tmp_path/'setup')
    html=output.read_text(encoding='utf-8')
    assert 'DO_NOT_EMBED' not in html
    assert 'Consultation' in html and 'M1' in html
    assert '"confirmed": false' in html


def test_local_status_override_does_not_require_sql_changes():
    profile=Profile(status_codes={'Locally finished':'completed'})
    assert profile.classify_status('Locally finished')=='completed'
    assert profile.classify_status('Pending')=='open'


@pytest.mark.parametrize('schema,permission,reason',[
    (1,0,'select_unavailable'),(0,0,'schema_unavailable'),(None,None,'unavailable_legacy'),
])
def test_profile_form_explains_missing_schema_or_read_permission(tmp_path,schema,permission,reason):
    wb=Workbook();wb.remove(wb.active)
    data={
        '00_Metadata':[['contract_version'],['2.0']],
        '01_Capabilities':[['contract_version','source_name','column_name','available','is_required','schema_available','select_allowed'],
                           ['2.0','DWH.Source','Time',0,0,schema,permission]],
        '02_Activities':[['activity_code','activity_name','activity_category'],['TX','Treatment','Clinical']],
    }
    for name,rows in data.items():
        sheet=wb.create_sheet(name)
        for row in rows:sheet.append(row)
    path=tmp_path/'source-check.xlsx';wb.save(path)
    html=create(path,tmp_path/'setup').read_text(encoding='utf-8')
    assert '"reason": "'+reason+'"' in html
    assert 'SELECT-Recht' in html


@pytest.mark.parametrize('existing',[False,True])
def test_profile_period_comes_from_export_without_manual_reentry(tmp_path,existing):
    wb=Workbook();wb.remove(wb.active)
    data={
        '00_Metadata':[['contract_version','period_start','period_end','site','period_reason'],
                       ['2.0',datetime(2024,4,1),datetime(2025,3,31),'Clinic-X','Geraetewechsel']],
        '01_Capabilities':[['contract_version','source_name','column_name','available','is_required'],['2.0','Source','Time',1,1]],
        '02_Activities':[['activity_code','activity_name','activity_category'],['TX','Treatment','Clinical']],
    }
    for name,rows in data.items():
        sheet=wb.create_sheet(name)
        for row in rows:sheet.append(row)
    path=tmp_path/'different-period.xlsx';wb.save(path)
    profile_path=tmp_path/'previous.json'
    profile_path.write_text(json.dumps(Profile(confirmed=True,activity_codes={'TX':'treatment_external'}).as_dict()),encoding='utf-8')
    html=create(path,tmp_path/'setup',profile_path if existing else None).read_text(encoding='utf-8')
    payload=json.loads(html.split('const DATA=',1)[1].split(', $=id',1)[0])
    profile=Profile(**payload['profile'])
    assert (profile.start,profile.end)==('2024-04-01','2025-03-31')
    assert profile.site=='Clinic-X' and profile.period_reason=='Geraetewechsel'
    assert not profile.confirmed and not profile.sources_complete
    if existing:assert profile.activity_codes=={'TX':'treatment_external'}
