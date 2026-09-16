"""Read an SSRS workbook once, never execute formulas or external workbook links."""
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook


def load_export(path: Path):
    if path.suffix.casefold()==".csv":
        events=pd.read_csv(path,dtype=str,keep_default_na=False,encoding="utf-8-sig").replace("",None)
        meta_fields={"contract_version","run_id","period_start","period_end","context_start",
                     "data_through","data_through_confirmed"}
        if events.empty or not meta_fields <= set(events):
            raise ValueError("Incomplete flat export metadata")
        meta_fields |= {"site", "period_reason", "comparison_population_complete", "collector_release"} & set(events)
        meta_rows=events[sorted(meta_fields)].drop_duplicates()
        if len(meta_rows)!=1 or str(meta_rows.iloc[0].contract_version)!="2.0":
            raise ValueError("Mixed or unsupported flat export contract")
        required={"source","event_key","patient_key","event_start","event_end","status","machine","activity_code"}
        if not required <= set(events):
            raise ValueError("Incomplete flat event contract")
        return meta_rows.iloc[0].to_dict(),events,[],[]
    tables = {"00_Metadata":[],"01_Capabilities":[],"02_Activities":[],"90_Events":[],"91_Images":[],
              "92_Acquisition":[]}
    with path.open("rb") as stream:
        workbook = load_workbook(stream,read_only=True,data_only=True,keep_links=False)
        try:
            for sheet in workbook:
                prefix = next((key for key in tables if sheet.title.startswith(key)),None)
                if not prefix:
                    continue
                headers = None
                for i,values in enumerate(sheet.iter_rows(values_only=True)):
                    if headers is None:
                        if "contract_version" in values or "activity_code" in values:
                            headers = [(j,str(v)) for j,v in enumerate(values) if v is not None]
                        elif i>40:
                            break
                        continue
                    record = {name:values[j] if j<len(values) else None for j,name in headers}
                    if not any(v is not None for v in record.values()):
                        continue
                    if record.get(headers[0][1]) == headers[0][1]:
                        continue
                    tables[prefix].append(record)
        finally:
            workbook.close()
    metadata = tables["00_Metadata"]
    if len(metadata)!=1 or str(metadata[0].get("contract_version")) not in {"2.0","2"}:
        raise ValueError("Unsupported export contract; collect with RDL 2.0")
    if not tables["90_Events"]:
        raise ValueError("No event details: image objects alone do not establish source coverage")
    events = pd.DataFrame(tables["90_Events"] + tables["91_Images"])
    if events.empty:
        raise ValueError("No event details: check collection_state, capabilities and export version")
    required = {"source","event_key","patient_key","event_start","event_end","status","machine","activity_code"}
    if not required <= set(events):
        raise ValueError("Incomplete event contract")
    if events.get("run_id",pd.Series(dtype=str)).nunique()!=1:
        raise ValueError("Mixed or missing run identity")
    if events.run_id.iloc[0] != metadata[0].get("run_id"):
        raise ValueError("Run identity differs from metadata")
    events = enrich_images(events, tables['92_Acquisition'], metadata[0])
    return metadata[0],events,tables["01_Capabilities"],tables["02_Activities"]


def enrich_images(events, records, metadata):
    """Require an exact same-run image key. Never manufacture a patient/day match."""
    if not records:
        metadata['image_acquisition_state'] = 'NOT_INCLUDED'
        return events
    native = pd.DataFrame(records)
    if set(native.run_id.dropna()) != {metadata['run_id']}:
        raise ValueError('Image acquisition run identity differs from metadata')
    metadata['image_acquisition_state'] = 'AVAILABLE' if native.source_state.eq('AVAILABLE').any() else 'UNAVAILABLE'
    if not events.source.eq('image_object').any():
        return events
    native = native[native.source_state.eq('AVAILABLE') & native.event_key.notna()].drop_duplicates()
    if native.event_key.duplicated().any():
        raise ValueError('Conflicting acquisition image identities')
    fields = ['event_key','acquisition_kind','image_manufacturer','acquisition_machine','acquisition_time']
    result = events.merge(native.reindex(columns=fields), on='event_key',how='left',validate='many_to_one')
    images = result.source.eq('image_object')
    known = images & result.acquisition_kind.notna() & result.acquisition_kind.ne('unknown')
    # Keep the more specific kV/MV label for manufacturer-qualified CBCT objects.
    keep_specific = result.acquisition_kind.eq('cbct_unknown') & result.image_kind.isin(['kv_cbct','mv_cbct'])
    result.loc[known & ~keep_specific,'image_kind'] = result.loc[known & ~keep_specific,'acquisition_kind']
    result.loc[known,'image_class_evidence'] = 'native_manufacturer_modality_reference'
    machine = images & result.acquisition_machine.fillna('').ne('')
    result.loc[machine,'machine'] = result.loc[machine,'acquisition_machine']
    return result


def normalize_flow(events, profile):
    frame = events.copy()
    frame["kind"] = profile.classify_activities(frame)
    delivered = frame.source.eq("delivery")
    brachy = pd.to_numeric(frame.get("is_brachy",pd.Series(0,index=frame.index)),errors="coerce").eq(1)
    frame.loc[delivered & ~brachy,"kind"] = "treatment_external"
    frame.loc[delivered & brachy,"kind"] = "treatment_brachy"
    # Calendar entries on R&V devices are scheduling evidence, never additional
    # clinical fractions, starts or episode endpoints. Keep them in throughput.
    connected = set(frame.loc[delivered, 'machine'].dropna()) - {''}
    scheduled_rv = (frame.source.eq('appointment') & frame.kind.str.startswith('treatment')
                    & frame.machine.isin(connected))
    frame = frame[~scheduled_rv]
    # Imaging is never a treatment-start event by itself.
    frame = frame[~frame.source.eq("imaging") & frame.patient_key.notna()]
    return frame


def eligible_events(events, profile):
    """Use non-identifying source flags when available; never guess from a hash."""
    frame = events.copy()
    if "patient_class" in frame:
        frame = frame[~frame.patient_class.eq("test_name")]
        if profile.require_numeric_patient_id:
            frame = frame[frame.patient_class.isin(["clinical_numeric", "no_patient", "unknown"])]
    if "resource_status" in frame:
        # A deleted/cancelled resource assignment is not a second clinical appointment.
        frame = frame[~frame.resource_status.fillna("").str.casefold().isin(["deleted", "cancelled"])
                      | frame.status.map(profile.classify_status).eq("cancelled")]
    return frame


def resolve_treatment_devices(events, profile):
    """Resolve missing/ambiguous external slots only from one actual device that day."""
    frame = events.copy()
    frame["machine"] = frame.machine.fillna("")
    date = pd.to_datetime(frame.event_start, errors="coerce", format="mixed").dt.strftime("%Y-%m-%d")
    technical = frame[frame.source.eq("delivery") & frame.machine.isin(profile.machines)].copy()
    technical["_day"] = date.loc[technical.index]
    candidates = technical.groupby(["patient_key", "_day"]).machine.agg(lambda values: tuple(set(values)))
    candidates = candidates[candidates.map(len).eq(1)].map(lambda values: values[0]).to_dict()
    selected = (frame.source.eq("appointment") & frame.machine.isin(["", "AMBIGUOUS_DEVICE"])
                & profile.classify_activities(frame).eq("treatment_external"))
    if "machine_inferred" not in frame:
        frame["machine_inferred"] = False
    for index in frame.index[selected]:
        machine = candidates.get((frame.at[index,"patient_key"], date.loc[index]))
        if machine:
            frame.at[index,"machine"] = machine
            frame.at[index,"machine_inferred"] = True
    return frame
