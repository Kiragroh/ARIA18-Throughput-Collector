"""Create a local offline profile form from the aggregate preflight workbook."""
import argparse
from datetime import datetime
import json
from pathlib import Path
import sys
import webbrowser
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from analysis.contracts import Profile


def read_preflight(path):
    tables={key:[] for key in ['Metadata','Capabilities','ActivityCatalog','AppointmentInventory','MachineInventory','HistoryStatusInventory']}
    wb=load_workbook(path,read_only=True,data_only=True,keep_links=False)
    try:
        for sheet in wb:
            key=next((k for k in tables if (k=='ActivityCatalog' and sheet.title.startswith('02_Activities')) or
                      sheet.title.split('_',1)[-1]==k),None)
            if not key:continue
            headers=None
            for cells in sheet.iter_rows(values_only=True):
                if headers is None:
                    if any(v in ('contract_version','activity_code','source_state') for v in cells):
                        headers=[(i,str(v)) for i,v in enumerate(cells) if v is not None]
                    continue
                row={h:cells[i] if i<len(cells) else None for i,h in headers}
                if any(v is not None for v in row.values()) and row.get(headers[0][1])!=headers[0][1]:tables[key].append(row)
    finally:wb.close()
    if not tables['Metadata'] or str(tables['Metadata'][0].get('contract_version')) not in {'2.0','2'}:
        raise ValueError('Preflight 2.0 required')
    if not tables['ActivityCatalog'] or not tables['Capabilities']:
        raise ValueError('Activity catalog and capabilities required')
    return tables


def create(path,output,existing=None):
    tables=read_preflight(path)
    profile=Profile().as_dict() if existing is None else Profile(**json.loads(existing.read_text(encoding='utf-8-sig'))).as_dict()
    profile['confirmed']=False;profile['sources_complete']=False;profile['complete_through']=''
    metadata=tables['Metadata'][0]
    for field,key in [('period_start','start'),('period_end','end')]:
        if metadata.get(field) is not None:
            profile[key]=datetime.fromisoformat(str(metadata[field])).date().isoformat()
    site=str(metadata.get('site') or '').strip()
    if site.casefold() not in {'','standort','\u00e4ndere mich','aendere mich'}:
        profile['site']=site
    if metadata.get('period_reason') is not None:
        profile['period_reason']=str(metadata['period_reason'])
    Profile(**profile)
    uses={}
    statuses=set()
    for row in tables['AppointmentInventory']:
        code=row.get('activity_code')
        if code:
            uses[code]=uses.get(code,0)+int(row.get('distinct_appointments') or 0)
        if row.get('appointment_status'):statuses.add(str(row['appointment_status']))
    activities=[dict(code=str(r['activity_code']),name=str(r.get('activity_name') or ''),
                     category=str(r.get('activity_category') or ''),count=uses.get(r['activity_code'],0))
                for r in tables['ActivityCatalog'] if r.get('activity_code')]
    machines=sorted({str(r['machine']) for key in ['MachineInventory','AppointmentInventory'] for r in tables[key]
                     if r.get('machine') and r['machine'] not in {'UNMAPPED','AMBIGUOUS_DEVICE'}})
    def flag(value):
        return str(value).casefold() in {'1','1.0','true'}
    missing=[]
    for row in tables['Capabilities']:
        if flag(row.get('available')):
            continue
        reason=('unavailable_legacy' if row.get('schema_available') is None or row.get('select_allowed') is None
                else 'schema_unavailable' if not flag(row['schema_available']) else 'select_unavailable')
        missing.append(dict(source=row.get('source_name'),column=row.get('column_name'),
                            required=flag(row.get('is_required')),reason=reason))
    # Only the allowlisted inventory is embedded, never arbitrary workbook cells.
    data=dict(profile=profile,activities=activities,machines=machines,statuses=sorted(statuses),missing=missing)
    payload=json.dumps(data,ensure_ascii=True).replace('<','\\u003c').replace('>','\\u003e').replace('&','\\u0026')
    template=(ROOT/'analysis/preflight_form.html').read_text(encoding='utf-8')
    output.mkdir(parents=True,exist_ok=True)
    target=output/'Standortprofil.html';target.write_text(template.replace('__SETUP_DATA__',payload),encoding='utf-8')
    print('PREFLIGHT_FORM_OK activities='+str(len(activities))+' machines='+str(len(machines)))
    return target


def main():
    parser=argparse.ArgumentParser();parser.add_argument('preflight',type=Path,nargs='?')
    parser.add_argument('--output',type=Path);parser.add_argument('--existing',type=Path);parser.add_argument('--open',action='store_true')
    args=parser.parse_args()
    interactive=args.preflight is None
    if interactive:
        from tkinter import Tk,filedialog
        root=Tk();root.withdraw()
        selected=filedialog.askopenfilename(title='Full Collector mit Prueftabellen als Excel auswaehlen',filetypes=[('Excel','*.xlsx')])
        root.destroy()
        if not selected:return
        args.preflight=Path(selected)
    target=create(args.preflight,args.output or args.preflight.parent/'Standort-Setup',args.existing)
    if args.open or interactive:webbrowser.open(target.resolve().as_uri())


if __name__=='__main__':main()
