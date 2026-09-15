"""Deterministic fabricated events. No clinical input is read."""
import argparse
from pathlib import Path
import sys
import random
from datetime import datetime,timedelta
from openpyxl import Workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from tools.build_collector_v2 import FIELDS


def create(output):
    rng=random.Random(2025)
    wb=Workbook(write_only=True)
    meta=wb.create_sheet("00_Metadata")
    meta.append(["contract_version","run_id","site","period_start","period_end","context_start","data_through","data_through_confirmed"])
    meta.append(["2.0","SYNTHETIC","Synthetisch",datetime(2025,1,1),datetime(2025,12,31),datetime(2024,1,1),datetime(2026,6,1),True])
    caps=wb.create_sheet("01_Capabilities")
    caps.append(["contract_version","source_name","column_name","is_required","available"])
    caps.append(["2.0","SYNTHETIC","SyntheticOnly",True,True])
    catalog=wb.create_sheet("02_Activities")
    catalog.append(["activity_code","activity_name","activity_category"])
    for code,name in [("TX","Therapie"),("CONS","Aufklaerung"),("RETURN","Wiedervorstellung"),("BREAK","Pause")]:
        catalog.append([code,name,"Synthetisch"])
    sheet=wb.create_sheet("90_Events")
    sheet.append(FIELDS)
    rows=[]

    def add(**row):
        record=dict(contract_version="2.0",run_id="SYNTHETIC",event_key="SYNTHETIC-"+str(len(rows)),
                    is_brachy=0,status="completed",time_source="measured",completion_candidates=1)
        record.update(row)
        rows.append(record)

    for month in range(1,13):
        for day in (6,7,8,9,10):
            for device in range(1,5):
                start=datetime(2025,month,day,8)
                for n in range(18):
                    patient=f"SIM-{month}-{day}-{device}-{n}"
                    duration=10+device+rng.uniform(-3,6)
                    beam=start+timedelta(minutes=5+device)
                    finish=beam+timedelta(minutes=duration-5)
                    end=finish+timedelta(minutes=3)
                    plan="PLAN-"+patient
                    if day==6:
                        add(source="appointment",patient_key=patient,event_start=start-timedelta(days=12),
                            activity_code="CONS",machine="",event_end=start-timedelta(days=12)+timedelta(minutes=30))
                        if n<6:
                            add(source="appointment",patient_key=patient,event_start=start-timedelta(days=5),
                                activity_code="CONS",machine="",event_end=start-timedelta(days=5)+timedelta(minutes=30))
                    add(source="delivery",patient_key=patient,machine=f"M{device}",event_start=beam,
                        event_end=finish,course_key="COURSE-"+patient,plan_key=plan,fraction=1,activity_code="")
                    add(source="imaging",patient_key=patient,machine=f"M{device}",event_start=beam-timedelta(minutes=2),
                        event_end=beam-timedelta(minutes=1),course_key="COURSE-"+patient,plan_key=plan,fraction=1,activity_code="")
                    add(source="appointment",patient_key=patient,machine=f"M{device}",event_start=start,
                        event_end=start+timedelta(minutes=20),activity_start=start+timedelta(minutes=3),
                        activity_end=end,completed=end+timedelta(minutes=1),activity_code="TX")
                    start=end+timedelta(minutes=35 if n==8 else rng.uniform(1,5))
    # Large enough strata to demonstrate reconciled flow without small cells.
    for n in range(30):
        add(source="appointment",patient_key="NO-"+str(n),machine="",event_start=datetime(2025,3,1,9)+timedelta(minutes=n),
            event_end=datetime(2025,3,1,9)+timedelta(minutes=n+30),activity_code="CONS")
        if n<10:
            add(source="appointment",patient_key="NO-"+str(n),machine="",event_start=datetime(2025,10,1,9),
                event_end=datetime(2025,10,1,10),activity_code="RETURN",status="open")
        add(source="appointment",patient_key="CAN-"+str(n),machine="",event_start=datetime(2025,2,1,9)+timedelta(minutes=n),
            event_end=datetime(2025,2,1,9)+timedelta(minutes=n+30),activity_code="CONS",status="cancelled")
    for row in rows:
        sheet.append([row.get(f) for f in FIELDS])
    # Exact resource fan-out duplicates are intentionally present.
    for row in rows[:50]:
        sheet.append([row.get(f) for f in FIELDS])
    output.parent.mkdir(parents=True,exist_ok=True)
    wb.save(output)
    print("SYNTHETIC_EXPORT_OK rows="+str(len(rows)+50))


if __name__=="__main__":
    parser=argparse.ArgumentParser()
    parser.add_argument("--output",type=Path,required=True)
    create(parser.parse_args().output)
